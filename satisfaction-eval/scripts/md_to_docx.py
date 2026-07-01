"""Convert the technical documentation markdown to a formatted Word document."""
import re
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

def set_cell_shading(cell, color):
    """Set cell background color."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def add_styled_paragraph(doc, text, style='Normal', bold=False, size=None, color=None, alignment=None):
    """Add a paragraph with custom formatting."""
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    if bold:
        run.bold = True
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    if alignment is not None:
        p.alignment = alignment
    return p

def parse_md_table(lines, start_idx):
    """Parse a markdown table starting at start_idx. Returns (rows, next_idx)."""
    header_line = lines[start_idx].strip()
    sep_line = lines[start_idx + 1].strip()
    if not sep_line.startswith('|') or '---' not in sep_line:
        return None, start_idx

    headers = [h.strip() for h in header_line.split('|')[1:-1]]
    rows = [headers]

    idx = start_idx + 2
    while idx < len(lines) and lines[idx].strip().startswith('|'):
        cells = [c.strip() for c in lines[idx].strip().split('|')[1:-1]]
        rows.append(cells)
        idx += 1
    return rows, idx

def convert_md_to_docx(md_path, docx_path):
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    # Configure heading styles
    for i in range(1, 4):
        heading_style = doc.styles[f'Heading {i}']
        hfont = heading_style.font
        hfont.name = '微软雅黑'
        heading_style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        hfont.color.rgb = RGBColor(0x1A, 0x56, 0xDB)  # Blue
        if i == 1:
            hfont.size = Pt(22)
        elif i == 2:
            hfont.size = Pt(16)
        else:
            hfont.size = Pt(13)

    # Set narrow margins
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # Skip empty lines
        if not line:
            i += 1
            continue

        # Headings
        if line.startswith('# ') and not line.startswith('## '):
            doc.add_heading(line[2:], level=1)
            i += 1
            continue
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=2)
            i += 1
            continue
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=3)
            i += 1
            continue
        elif line.startswith('#### '):
            doc.add_heading(line[5:], level=4)
            i += 1
            continue

        # Horizontal rule
        if line.strip() == '---':
            doc.add_paragraph('─' * 60)
            i += 1
            continue

        # Code blocks (```...```)
        if line.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code_lines.append(lines[i].rstrip())
                i += 1
            i += 1  # skip closing ```
            if code_lines:
                code_text = '\n'.join(code_lines)
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1)
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                # Add light gray background via shading
                pPr = p._p.get_or_add_pPr()
                shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F5F5F5"/>')
                pPr.append(shading)
            continue

        # Tables
        if line.startswith('|') and i + 2 < len(lines):
            table_data, next_idx = parse_md_table(lines, i)
            if table_data:
                headers = table_data[0]
                data = table_data[1:]
                table = doc.add_table(rows=len(table_data), cols=len(headers))
                table.style = 'Light Grid Accent 1'
                table.alignment = WD_TABLE_ALIGNMENT.CENTER

                # Header row
                for j, h in enumerate(headers):
                    cell = table.rows[0].cells[j]
                    cell.text = h
                    for para in cell.paragraphs:
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in para.runs:
                            run.bold = True
                            run.font.size = Pt(10)
                    set_cell_shading(cell, '1A56DB')
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

                # Data rows
                for r_idx, row in enumerate(data):
                    for c_idx, val in enumerate(row):
                        cell = table.rows[r_idx + 1].cells[c_idx]
                        cell.text = val
                        for para in cell.paragraphs:
                            for run in para.runs:
                                run.font.size = Pt(9.5)

                i = next_idx
                doc.add_paragraph()  # spacing after table
                continue

        # Inline code `text`
        if '`' in line:
            parts = re.split(r'(`[^`]+`)', line)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            for part in parts:
                if part.startswith('`') and part.endswith('`'):
                    run = p.add_run(part[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
                else:
                    # Handle bold and other markers
                    clean = part.replace('**', '')
                    run = p.add_run(clean)
                    run.font.size = Pt(11)
                    if part.count('**') >= 2:
                        run.bold = True
            i += 1
            continue

        # Regular text
        clean = line
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)

        # Handle bold markers
        if '**' in clean:
            parts = re.split(r'(\*\*[^*]+\*\*)', clean)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                    run.font.size = Pt(11)
                else:
                    run = p.add_run(part)
                    run.font.size = Pt(11)
        else:
            run = p.add_run(clean)
            run.font.size = Pt(11)

        i += 1

    # Add line spacing
    for para in doc.paragraphs:
        para.paragraph_format.line_spacing = Pt(20)

    doc.save(docx_path)
    print(f"Word document saved to: {docx_path}")

if __name__ == "__main__":
    md_path = "D:/claude code test/满意度评价/满意度评价运营平台-技术文档.md"
    docx_path = "D:/claude code test/满意度评价/满意度评价运营平台-技术文档.docx"
    convert_md_to_docx(md_path, docx_path)
